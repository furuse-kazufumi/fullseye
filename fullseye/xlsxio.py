# Copyright (c) 2026 Kazufumi Furuse. Licensed under the Apache License, Version 2.0 (see LICENSE).
"""xlsxio — 型付きの結果を Excel(.xlsx)レポートに書き出す(jsonio / mdio の第 3 系統)。

jsonio が「機械が bit で戻せる JSON」、mdio が「人が読む Markdown」を出すのに対し、こちらは
**現場が使う .xlsx** を出す —— 測定表・数値・スカラをセルに、画像はサムネイルを埋め込む。
検査レポートを Excel で配る/貼る用途に直結する。``sections`` は mdio.report と同じ
``(見出し, value, sort)`` の列なので、同じ材料から Markdown も JSON も Excel も出せる。

openpyxl(optional 依存: ``pip install "fullseye[xlsx]"``)を使う。無い環境では明示的に
ImportError にする(黙って別形式にしない)。画像の埋め込みに失敗しても**レポート自体は
壊さない**(その節は形・値域の 1 行要約セルに落とす、fail-soft)。
"""
from __future__ import annotations

import os
import tempfile

import numpy as np

from fullseye import jsonio as _J

#: セルに素で書ける型付き結果。画像・場・ボリュームはサムネイル埋め込みか要約に回す。
_TABULAR = frozenset({"table", "points", "matrix", "signal", "vector", "keypoints", "counts"})
_SCALARISH = frozenset({"feature", "scalar"})
_IMAGELIKE = frozenset({"image", "image2d", "color", "rgb", "rgbimage", "mask", "region"})


def _require_openpyxl():
    try:
        import openpyxl                                        # noqa: F401
        from openpyxl.drawing.image import Image as XLImage    # noqa: F401
        from openpyxl.styles import Font                       # noqa: F401
    except Exception as exc:                                    # noqa: BLE001
        raise ImportError(
            "xlsxio needs openpyxl — install with: pip install \"fullseye[xlsx]\" "
            "(or `pip install openpyxl`)") from exc
    return openpyxl


def _thumb(arr, longest=256):
    """大きい画像をストライドで縮めた [0,1] のサムネイル(PIL 不要)。"""
    a = np.asarray(arr, np.float64)
    if a.ndim < 2:
        return None
    h, w = a.shape[:2]
    step = max(1, int(max(h, w) / longest))
    a = a[::step, ::step]
    return np.clip(a, 0.0, 1.0)


def _write_tabular(ws, r, value, sort, Font, float_fmt, max_rows):
    """table / points / matrix / signal … をセルに。次に書ける行 r を返す。"""
    if sort == "table":
        rows = list(value) if not isinstance(value, np.ndarray) else [
            {"c%d" % j: v for j, v in enumerate(row)} for row in np.atleast_2d(value)]
        keys: list = []
        for d in rows:
            for k in (d.keys() if isinstance(d, dict) else range(len(d))):
                if k not in keys:
                    keys.append(k)
        for j, k in enumerate(keys, 1):
            ws.cell(r, j, str(k)).font = Font(bold=True)
        r += 1
        for d in rows[:max_rows]:
            for j, k in enumerate(keys, 1):
                v = d.get(k) if isinstance(d, dict) else (d[k] if k < len(d) else None)
                ws.cell(r, j, v if isinstance(v, (int, float, str)) or v is None else str(v))
            r += 1
        if len(rows) > max_rows:
            ws.cell(r, 1, "… %d 行中 %d 行のみ表示" % (len(rows), max_rows))
            r += 1
        return r + 1
    a = np.atleast_2d(np.asarray(value, np.float64))
    if a.shape[0] == 1 and sort in ("signal", "vector"):
        a = a.reshape(-1, 1)                                    # 1-D は縦に並べる
    for i in range(min(a.shape[0], max_rows)):
        for j in range(a.shape[1]):
            ws.cell(r, j + 1, float(a[i, j]))
        r += 1
    if a.shape[0] > max_rows:
        ws.cell(r, 1, "… %d 行中 %d 行のみ表示" % (a.shape[0], max_rows))
        r += 1
    return r + 1


def save_xlsx_report(sections, path, *, title=None, float_fmt="%.6g", max_rows=1000,
                     thumbnails=True):
    """``sections`` を 1 枚の Excel シートに縦に並べて ``path`` に書き出す。

    ``sections`` は ``(見出し, value, sort)`` の列(mdio.report と同じ)。``table`` /
    ``points`` / ``matrix`` / ``signal`` / ``vector`` / ``keypoints`` / ``counts`` は
    セルの表に、``feature`` / ``scalar`` は 1 セルに、画像系(``image`` / ``color`` /
    ``region`` …)は ``thumbnails=True`` ならサムネイル PNG を埋め込む(失敗しても
    形・値域の 1 行要約に落とす=fail-soft)。書き出した ``path`` を返す。

    openpyxl が要る(``pip install "fullseye[xlsx]"``)。未知 sort は ``ValueError``。"""
    openpyxl = _require_openpyxl()
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font

    for s in sections:
        if len(s) != 3:
            raise ValueError("each section must be (heading, value, sort), got %r" % (s,))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    r = 1
    if title:
        c = ws.cell(r, 1, title)
        c.font = Font(bold=True, size=14)
        r += 2
    tmpfiles: list[str] = []
    tmpdir = tempfile.mkdtemp(prefix="fs_xlsx_")
    try:
        from api import write_image
        for heading, value, sort in sections:
            hc = ws.cell(r, 1, str(heading))
            hc.font = Font(bold=True, size=12)
            r += 1
            if sort in _TABULAR:
                r = _write_tabular(ws, r, value, sort, Font, float_fmt, max_rows)
            elif sort in _SCALARISH:
                v = _J.as_value(value) if _J.is_envelope(value) else value
                ws.cell(r, 1, float(v) if isinstance(v, (int, float, np.floating)) else str(v))
                r += 2
            elif sort in _IMAGELIKE:
                placed = False
                if thumbnails:
                    try:
                        th = _thumb(value)
                        if th is not None:
                            p = os.path.join(tmpdir, "img_%d.png" % r)
                            write_image(p, th)
                            tmpfiles.append(p)
                            ws.add_image(XLImage(p), ws.cell(r, 1).coordinate)
                            r += max(8, min(30, th.shape[0] // 20))
                            placed = True
                    except Exception:                          # noqa: BLE001 (fail-soft)
                        placed = False
                if not placed:
                    a = np.asarray(value, np.float64)
                    ws.cell(r, 1, "image %s  range [%.4g, %.4g]" % (
                        tuple(a.shape), float(a.min()) if a.size else 0.0,
                        float(a.max()) if a.size else 0.0))
                    r += 2
            else:
                raise ValueError("xlsxio: unknown sort %r (heading %r)" % (sort, heading))
        wb.save(path)
    finally:
        for p in tmpfiles:
            try:
                os.remove(p)
            except OSError:
                pass
        try:
            os.rmdir(tmpdir)
        except OSError:
            pass
    return path


__all__ = ["save_xlsx_report"]
