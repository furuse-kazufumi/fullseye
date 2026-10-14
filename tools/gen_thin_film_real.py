# Copyright (c) 2026 Kazufumi Furuse. Licensed under the Apache License, Version 2.0 (see LICENSE).
"""薄膜の PoC の実測データ(CC BY 4.0 の 2 件)を、取得した原本から ``data/thin_film_real/`` の小さな表に抜き出す。

    py -3.11 tools/gen_thin_film_real.py <取得した原本のディレクトリ>

原本(このスクリプトは取得しない。手で落として 1 つのディレクトリに置く):
  * Zenodo 10.5281/zenodo.13753754(CC BY 4.0): Schwarz ら「Spectro-Spatial Unmixing in Optical Microspectroscopy
    for Thickness Determination of Layered Materials」Adv. Opt. Mater. (2024), doi:10.1002/adom.202402502 の
    データ。``Supporting_Information_Figure_Data.zip`` の ``Figure_S3a_data.xlsx``(SiO₂/Si の実測反射率、
    被覆率 11 段)、``Manuscript_Figure_Data.zip`` の ``Figure_4_data.xlsx``(t_Model / t_Reference)、
    ``Micrographs_SiO2_Sample.zip`` の ``48.4%_SiO2_Sample.jpg``(顕微鏡のカラー写真)。
    → ``<src>/zenodo_13753754/ex/`` に展開して置く。
  * figshare(Sheffield ORDA)10.15131/shef.data.23285603(CC BY 4.0): Sánchez-Arriaga ら「A Spectroscopic
    Reflectance-Based Low-Cost Thickness Measurement System for Thin Films」Sensors 23, 5326 (2023),
    doi:10.3390/s23115326 のデータ。``CFM_DATA_*.xlsx`` の 3 本(市販の SiO₂/Si 膜厚標準 476.3 nm・198.7 nm の実測反射、繰り返し 3〜4 回)。
    装置の型番は公開側に書かない(ファイル名は sensor1 / sensor2 / sensor2_lamp に付け替える)。
    → ``<src>/sheffield_23285603/`` に置く。

改変(CC BY の表示義務として ``data/thin_film_real/README.md`` に書く): 表から必要な列だけを抜いた。Sheffield は
繰り返しの平均を取った。写真は 1920×1080 から 640×360 へ縮小し JPEG で保存し直した。原本は untrusted として
読むだけ(openpyxl の read_only・PIL)で、何も実行しない。
"""
from __future__ import annotations

import csv
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "thin_film_real")


def _rows(path, sheet=None):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return list(ws.iter_rows(values_only=True))


def zenodo(src):
    ex = os.path.join(src, "zenodo_13753754", "ex")
    rows = _rows(os.path.join(ex, "Figure_S3a_data.xlsx"))
    cov = [float(x) for x in rows[1][2:] if x is not None]
    with open(os.path.join(OUT, "schwarz2024_sio2_si_reflectance.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["wavelength_nm"] + ["R_cov%.2f" % c for c in cov])
        for r in rows[2:]:
            if r[1] is None:
                continue
            w.writerow(["%.4f" % float(r[1])] + ["%.6f" % float(v) for v in r[2:2 + len(cov)]])
    r4 = _rows(os.path.join(ex, "Figure_4_data.xlsx"))
    with open(os.path.join(OUT, "schwarz2024_t_model_over_reference.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["sio2_coverage_percent", "t_model_over_t_reference"])
        for r in r4[1:]:
            if r[0] is not None and r[1] is not None:
                w.writerow(["%.2f" % float(r[0]), "%.6f" % float(r[1])])
    from PIL import Image
    im = Image.open(os.path.join(ex, "Micrographs_SiO2_Sample", "48.4%_SiO2_Sample.jpg")).convert("RGB")
    im.resize((640, 360), Image.LANCZOS).save(os.path.join(OUT, "schwarz2024_micrograph_sio2_edge_640.jpg"),
                                              quality=92)


def sheffield(src):
    d = os.path.join(src, "sheffield_23285603")
    out = []
    import glob
    names = sorted(os.path.basename(f) for f in glob.glob(os.path.join(d, "CFM_DATA_*.xlsx")))
    labels = {}
    for fn in names:
        k = "sensor1" if "SENSOR1" in fn else "sensor2"
        labels[fn] = k + ("_lamp" if "HAL" in fn else "")
    for fn in names:
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(d, fn), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rr = list(ws.iter_rows(values_only=True))
            hdr = rr[2]
            tcols = [i for i, h in enumerate(hdr) if isinstance(h, str) and h.strip().lower().startswith("test")]
            sample = 1 if "SAMPLE1" in ws.title.replace(" ", "") else 2
            for r in rr[3:]:
                if not isinstance(r[0], (int, float)):
                    continue
                vals = [r[i] for i in tcols if isinstance(r[i], (int, float))]
                if len(vals) != len(tcols):
                    continue
                out.append((labels[fn], ws.title.strip(), sample, float(r[0]), sum(vals) / len(vals), len(vals)))
    with open(os.path.join(OUT, "sanchez_arriaga2023_sio2_standards.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["file", "sheet", "sample", "wavelength_nm", "R_mean_arb", "n_repeats"])
        for o in out:
            w.writerow([o[0], o[1], o[2], "%.3f" % o[3], "%.6f" % o[4], o[5]])


def main(argv):
    if len(argv) != 2 or not os.path.isdir(argv[1]):
        print(__doc__)
        return 2
    os.makedirs(OUT, exist_ok=True)
    zenodo(argv[1])
    sheffield(argv[1])
    print("wrote", OUT)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
