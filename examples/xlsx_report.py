# Copyright (c) 2026 Kazufumi Furuse. Licensed under the Apache License, Version 2.0 (see LICENSE).
"""xlsx_report — 型付きの検査結果を Excel(.xlsx)レポートに書き出す(jsonio/mdio の第3系統)。

    py -3.11 examples/xlsx_report.py

【この例が示すこと】
mdio.report と同じ ``(見出し, value, sort)`` の列から、現場が使う .xlsx を作る。測定表・
点群・スカラはセルに、画像はサムネイルを 1 枚埋め込む。書いたファイルを openpyxl で開き直し、
セルの値と埋め込み画像の枚数を assert で確かめる(絵に描いたふりをしない)。

EXTEND: 同じ ``sections`` を fs.report(Markdown)/ fs.save_json(機械)にも渡せる。
openpyxl が無い環境では ``pip install "fullseye[xlsx]"``。この例はその場合 skip 表示で PASS。
"""
from __future__ import annotations

import os
import sys
import tempfile

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import fullseye as fs  # noqa: E402


def main() -> int:
    try:
        import openpyxl
    except ImportError:
        print('SKIP: openpyxl 未導入(pip install "fullseye[xlsx]")。PASS(skip)')
        return 0

    # 検査結果に見立てた材料 —— blob 表 / 点群 / スカラ特徴 / プレビュー画像。
    blobs = [{"id": 1, "area": 12.5, "label": "OK"},
             {"id": 2, "area": 3.0, "label": "NG"}]
    pts = np.array([[10.5, 20.0], [30.25, 40.0], [50.0, 60.5]])
    yy, xx = np.mgrid[0:160, 0:160]
    preview = 0.5 + 0.5 * np.sin(xx / 18.0) * np.cos(yy / 22.0)

    sections = [("欠陥ブロブ", blobs, "table"),
                ("重心座標", pts, "points"),
                ("被覆率", 0.42, "feature"),
                ("プレビュー", preview, "image")]

    out = os.path.join(tempfile.gettempdir(), "fullseye_inspection_report.xlsx")
    path = fs.save_xlsx_report(sections, out, title="検査レポート")
    print("書き出し:", path)

    # 開き直して検算 —— セルの値と埋め込み画像。
    wb = openpyxl.load_workbook(path)
    ws = wb["Report"]
    vals = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert ws["A1"].value == "検査レポート"
    assert 12.5 in vals and "label" in vals, "blob 表がセルに出ていない"
    assert 0.42 in vals, "スカラ特徴がセルに出ていない"
    assert 10.5 in vals and 60.5 in vals, "点群がセルに出ていない"
    assert len(ws._images) == 1, "プレビュー画像が埋め込まれていない"
    print("表・点群・スカラ・埋め込み画像 1 枚 を確認")

    # 同じ材料が Markdown にもなる(第3系統は出力先が違うだけ)。
    md = fs.report(sections[:3], title="検査レポート")
    assert "欠陥ブロブ" in md and "被覆率" in md
    print("同じ sections から Markdown も生成: OK")
    print("PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
